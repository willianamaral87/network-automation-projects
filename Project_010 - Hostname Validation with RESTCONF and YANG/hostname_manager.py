# Hostname tool for Cisco IOS-XE switches (RESTCONF)
# Menu: 1 collect | 2 validate | 3 change | 4 save | 5 exit

# Functions from helpers.py:
# read_excel -> reads inventory.xlsx and returns the list of devices
# save_excel -> saves the IP/hostname list to an Excel file
from helpers import read_excel, save_excel

# RESTCONF functions: GET reads data, PATCH changes data
from restconf import restconf_get, restconf_patch

# pauses the script (time.sleep)
import time 


import subprocess   # runs external commands (yanglint) from Python
import requests     # sends HTTP/RESTCONF requests

# writes the payload to a .json file
import json

from openpyxl import load_workbook   # reads Excel files (install: pip install openpyxl)

# Use getpass to read the credencial from user
from getpass import getpass
 
# Hide the "InsecureRequestWarning" caused by verify=False (self-signed cert)
requests.packages.urllib3.disable_warnings()

# RESTCONF path to the hostname (Cisco IOS-XE native model)
HOSTNAME_PATH = "/restconf/data/Cisco-IOS-XE-native:native/hostname"


# Ask for the device credentials (the password is not shown while typing)
def get_credentials():
    user = input("Username: ")
    pwd = getpass("Password: ")
    return user, pwd

# ---------------- Option 2.2: validate the CURRENT hostnames ----
# Reads the file created in option 1 and checks each hostname with yanglint.
# Returns how many hostnames are invalid (0 = all OK).
def validate_hostname_from_exported_excel():
    # Counts how many hostnames are invalid
    var_control = 0

    # Open the Excel file (column A = IP, column B = hostname, row 1 = header)
    sheet = load_workbook("hostname_collected_from_devices.xlsx").active

    # Read each row, starting at row 2 (skip the header), only columns A and B

    for host, hostname in sheet.iter_rows(min_row=2, max_col=2, values_only=True):
        
        # Build the payload with this hostname and save it to payload.json
        # (yanglint reads the data from this file)
        payload = {"Cisco-IOS-XE-native:native": {"hostname": hostname}}

        with open("payload.json", "w") as f:
            json.dump(payload, f)

        # Validate the payload with yanglint:
        #   -p xe/17131/  -> folder where yanglint finds the YANG models
        #   -t edit       -> the payload is a partial config (edit), not the full config
        #   Cisco-IOS-XE-native.yang         -> Cisco model
        #   company-hostname-deviation.yang  -> company hostname rule
        result = subprocess.run(
            ["yanglint", "-Q", "-t", "edit", "-p", "xe/17131/",
            "xe/17131/Cisco-IOS-XE-native.yang",
            "xe/17131/company-hostname-deviation.yang",
            "payload.json"],
            capture_output=True, text=True,
        )

        # yanglint exit code: 0 = valid, anything else = invalid
        if result.returncode == 0:
            print(f"OK       {host}  {hostname}")    
        else:
            print(f"INVALID  {host}  {hostname}  -> {result.stderr.splitlines()[0]}") # Show only the first line of the yanglint error

            # One more invalid hostname
            var_control +=1

    return var_control

# ---------------- Option 2.1: validate the NEW hostnames --------
# Reads inventory.xlsx and checks each new hostname with yanglint.
# Also used by option 3 before changing anything.
# Returns how many hostnames are invalid (0 = all OK).
def validate_new_hostname_from_inventory():

    # Open the Excel file (column A = IP, column B = new hostname, row 1 = header)
    sheet = load_workbook("inventory.xlsx").active

    # Counts how many hostnames are invalid
    var_control = 0

    # Read each row of the sheet:
    #   min_row=2          -> start at row 2 (row 1 is the header)
    #   max_col=2          -> read only columns A and B
    #   values_only=True   -> return the cell values, not cell objects
    #   host, new_hostname -> column A goes to host, column B to new_hostname
    for host, new_hostname in sheet.iter_rows(min_row=2, max_col=2, values_only=True):

        # Build the payload with this hostname and save it to payload.json
        # (yanglint reads the data from this file)
        payload = {"Cisco-IOS-XE-native:native": {"hostname": new_hostname}}

        with open("payload.json", "w") as f:
            json.dump(payload, f)

        # Validate with yanglint (Cisco model + company hostname rule)
        # Same command as in validate_hostname_from_exported_excel()
        result = subprocess.run(
            ["yanglint", "-Q", "-t", "edit", "-p", "xe/17131/",
            "xe/17131/Cisco-IOS-XE-native.yang",
            "xe/17131/company-hostname-deviation.yang",
            "payload.json"],
            capture_output=True, text=True,
        )

        # yanglint exit code: 0 = valid, anything else = invalid
        if result.returncode == 0:
            print(f"OK       {host}  {new_hostname}")
        else:
            print(f"INVALID  {host}  {new_hostname}  -> {result.stderr.splitlines()[0]}")

            # One more invalid hostname
            var_control = var_control+1

    return var_control


# ---------------- Main menu --------------------------------------
def main():
    print("\n===== HOSTNAME MENU =====")
    print("1. Collect the hostname on each switch" )
    print("2. Validate hostname from Excel file")
    print("3. Change hostname")
    print("4. Save configuration")
    print("5. Exit")

    # Keep asking until the user types a number from 1 to 5
    while True:
        option = input("\nSelect an option (1-5): ")

        if option.isdigit() and 1 <= int(option) <= 5:
            option = int(option)
            break
        else:
            print("Invalid option. Please enter a number between 1 and 5.")

    # Stores [IP, hostname] for each switch (exported to Excel in option 1)
    output_export = []

    # Read inventory.xlsx: one dictionary per row,
    # e.g. {"IP": "192.168.1.13", "new_hostname": "CRZ-SW-IT-25"}
    #
    devices = read_excel("inventory.xlsx")

    # Only options 1 and 3 connect to the switches
    if option in (1, 3):
        #user, pwd = get_credentials()
        user = "willian.amaral.ti"
        pwd = "6wzpY4O4__TLz"

    # OPTION 1 - COLLECT THE HOSTNAME
    # Reads the current hostname of each switch listed in inventory.xlsx.
    # Column A must have the IP address.
    # Column B has the new hostname (only needed if the hostname will be changed).
    if option == 1:
        print("Collecting the hostname from each switch....")

        for device in devices:
            IP = device['IP']

            ##

            # WARNING: this PATCH CHANGES the hostname on the switch.
            # It is here only for lab tests. Option 1 should only collect the hostname.
            
            #------------------------------
                    # Since I'm testing this script with a single device, this section reads the
                    # inventory file and changes the hostname on each iteration to simulate
                    # multiple devices on the network.
                    # If you're using more than one device, comment out the 3 lines below.

            new_hostname = device["new_hostname"]
            payload = {"Cisco-IOS-XE-native:hostname": new_hostname}
            restconf_patch(IP, user, pwd, HOSTNAME_PATH, payload)
            #------------------------------
            
            # Read the current hostname from the switch (RESTCONF GET)
            data = restconf_get(IP, user, pwd, HOSTNAME_PATH)

            # The answer uses the Cisco-IOS-XE-native model, e.g.
            #   {"Cisco-IOS-XE-native:hostname": "SW-01"}
            # Take only the hostname value
            current = data["Cisco-IOS-XE-native:hostname"]

            print(f"IP: {IP} - Hostname: {current} ")

            # Add [IP, hostname] to the list; it is used to create the Excel file
            output_export.append([IP, current])


            time.sleep(7)

        # Ask if the user wants to export the result to Excel
        option_1 = input("\nWould you like to export the output to Excel? (Y to confirm) ")

        # Y/y (yes) or S/s (sim) = export the result to Excel
        if option_1 == "Y" or option_1 == "y" or option_1 == "S" or option_1 == "s": 
            save_excel(output_export)
            
    # OPTION 2 - VALIDATE HOSTNAMES WITH YANGLINT
    # Two possible sources:
    #   1 - inventory.xlsx: the NEW hostnames, before they are configured
    #   2 - hostname_collected_from_devices.xlsx (created in option 1):
    #       the CURRENT hostnames on the switches
    elif option == 2:
        option_validate_hostname = int(input("Select: \n1 - Inventory file with the new hostnames, which will be used to configure the new devices. "+
                                             "\n2 - File generated in Step 1 to validate the current hostnames of the network switches.\n"))

        # Source 1: inventory.xlsx (new hostnames)
        if option_validate_hostname == 1:
            validate_new_hostname_from_inventory()
        # Source 2: file created in option 1 (current hostnames)
        elif option_validate_hostname == 2:
            print("Validating the current hostname from switches")

            validate_hostname_from_exported_excel() 
        else:
            print("Invalid option")

    # OPTION 3 - CHANGE THE HOSTNAME
    # First validate ALL new hostnames in inventory.xlsx.
    # If at least one is invalid, stop and change nothing.
    elif option == 3:
        print("Function changing the hostname")

        if validate_new_hostname_from_inventory() >= 1 :
            print("There is an error in the hostname...\nAborting the script. Nothing was changed on the device")
            exit(1)
        else:
            print('Sending config to each device')
       
            for device in devices:
                IP = device['IP']

                # Wait X second between devices
                time.sleep(7)

                # Send the new hostname to the switch (RESTCONF PATCH)
                new_hostname = device["new_hostname"]
                payload = {"Cisco-IOS-XE-native:hostname": new_hostname}
                restconf_patch(IP ,user,pwd,HOSTNAME_PATH,payload)
                   
                # Read the hostname again to confirm the change
                data = restconf_get(IP, user, pwd, HOSTNAME_PATH)

                current = data["Cisco-IOS-XE-native:hostname"]
                print(f"IP: {IP} - Hostname: {current} ")

    # OPTION 4 - SAVE THE CONFIGURATION (not implemented yet)
    elif option == 4:
        print("save_config")

    # OPTION 5 - EXIT
    elif option == 5:
        print("Exiting...")
        exit(1)
    else:
        print("Invalid option. Please run the script again.")

# ---------------- Start the program -----------------------------
# Python sets __name__ automatically:
#   "__main__"        -> you ran this file directly (python script.py)
#   something else    -> another script imported this file
# So main() runs ONLY when you run this file yourself.
if __name__ == "__main__":
    main()



