# reads YAML files (group_vars)
import yaml

# date and time, used in file names
from datetime import datetime

# reads Excel files
from openpyxl import load_workbook

# Read a YAML file and return its content
def read_yaml(path):
    with open(path, 'r') as f:
        device_list = yaml.safe_load(f)
    return device_list

# Write data to a file (replaces the old content)
def save_file(path, data):
    with open(path,'w') as f:
        f.write(data)

# Add data to the end of a file (keeps the old content)
def save_file_append(path, data):
    with open(path,'a') as f:
        f.write(data)

# Save data to a new file in the output/ folder, with date and time in the name
# e.g. output/output_20261013_084900.txt
def save_file_append_with_datetime(data):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"output_{timestamp}.txt"
    path = f"output/{filename}"

    with open(path, "a") as f:
        f.write(data)

# Read an Excel file and return a list of dicts, one per row
def read_excel(path):
    # data_only=True gives cell VALUES, not formulas
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active

    # All rows as tuples, e.g. ("192.168.1.13", "VAR-SW-IT-25")
    rows = list(sheet.iter_rows(values_only=True))

    # First row = column names -> become the dictionary keys
    headers = rows[0]

    # List of devices to return
    devices = []
    for row in rows[1:]:          # skip the header row
        if not any(row):          # skip empty rows
            continue
        # e.g. {"IP": "192.168.1.13", "new_hostname": "VAR-SW-IT-25"}
        devices.append(dict(zip(headers, row)))

    return devices

# Save the output to Excel File
def save_excel(output_export):
    # creates new Excel files
    from openpyxl import Workbook

    wb = Workbook()                     # create a new Excel file
    ws = wb.active                      # get the first sheet
    ws.title = "Hostnames"

    ws.append(["IP", "Hostname"])       # header row

    for row in output_export:           # each row is ['192.168.101.12', 'CSR1000V3']
        #print(f"row: {row}")
        ws.append(row)

    wb.save("hostname_collected_from_devices.xlsx")

    print(f"The name of the file is:")
    print("hostname_collected_from_devices.xlsx")